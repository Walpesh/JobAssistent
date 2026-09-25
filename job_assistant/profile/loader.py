"""
Загрузка профиля пользователя из DOCX.
"""

import os
from docx import Document
from job_assistant.utils.logging import add_log


def load_profile_from_docx(file_path: str) -> str | None:
    """
    Прочитать текст профиля из .docx (параграфы + таблицы).
    Возвращает None, если файл не найден или произошла ошибка.
    """
    if not os.path.exists(file_path):
        add_log("WARNING", f"Файл профиля {file_path} не найден.")
        return None
    try:
        doc = Document(file_path)
        full_text = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        full_text.append(cell.text)
        add_log("SUCCESS", f"Профиль успешно загружен из {file_path}")
        return "\n".join(full_text)
    except Exception as e:
        add_log("ERROR", f"Ошибка чтения файла профиля: {e}")
        return None
